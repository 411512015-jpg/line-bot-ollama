"""把報告與 AI 摘要轉成 Word（.docx），方便再編輯。"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"

GREEN = RGBColor(0x04, 0x9B, 0x43)
DARK = RGBColor(0x12, 0x20, 0x19)
MUTED = RGBColor(0x66, 0x75, 0x6C)
FONT = "Microsoft JhengHei"
MONO = "Consolas"

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def set_run_font(run, name=FONT, size=10.5, bold=None, color=None, italic=None):
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def shade(element, fill: str):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    element.get_or_add_pPr().append(shd)


def shade_cell(cell, fill: str):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def add_inline(paragraph, text, size=10.5, color=DARK):
    for token in INLINE.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True, color=RGBColor(0x07, 0x5F, 0x31))
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name=MONO, size=size - 1.2, color=RGBColor(0x17, 0x4D, 0x2C))
        elif token.startswith("[") and "](" in token:
            label, url = token[1:-1].split("](", 1)
            run = paragraph.add_run(f"{label}（{url}）")
            set_run_font(run, size=size, color=RGBColor(0x05, 0x7A, 0x39))
            run.font.underline = True
        else:
            run = paragraph.add_run(token)
            set_run_font(run, size=size, color=color)


def body_para(doc, text, style=None, size=10.5):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.32
    add_inline(p, text, size=size)
    return p


def heading(doc, text, level):
    sizes = {1: 24, 2: 15.5, 3: 12}
    colors = {1: DARK, 2: GREEN, 3: DARK}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == 2 else 12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, size=sizes.get(level, 12), bold=True, color=colors.get(level, DARK))
    if level == 2:
        pPr = p._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), "BFE9CB")
        pbdr.append(bottom)
        pPr.append(pbdr)
    return p


def add_table(doc, rows):
    cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            add_inline(p, value, size=9)
            if i == 0:
                shade_cell(cell, "E4F7E9")
                for run in p.runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(0x07, 0x5F, 0x31)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_code(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.left_indent = Cm(0.3)
    shade(p._p, "14251C")
    for i, line in enumerate(lines):
        run = p.add_run(line + ("\n" if i < len(lines) - 1 else ""))
        set_run_font(run, name=MONO, size=8.8, color=RGBColor(0xDF, 0xF8, 0xE7))


def add_quote(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(7)
    shade(p._p, "F0F8F2")
    add_inline(p, " ".join(line.lstrip("> ").strip() for line in lines), size=9.6,
               color=RGBColor(0x31, 0x53, 0x43))


def add_image(doc, path):
    doc.add_picture(str(path), width=Cm(16.2))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.space_before = Pt(6)
    doc.paragraphs[-1].paragraph_format.space_after = Pt(8)


def add_page_number_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("LINE Bot × Local AI｜")
    set_run_font(run, size=8, color=RGBColor(0x8B, 0x9A, 0x91))
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    set_run_font(run, size=8, color=RGBColor(0x8B, 0x9A, 0x91))


def setup(doc, cover=False):
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    add_page_number_footer(section)

    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def add_cover(doc, lines):
    for _ in range(5):
        doc.add_paragraph()
    label = doc.add_paragraph()
    r = label.add_run("MIDTERM PROJECT / 2026")
    set_run_font(r, size=8.5, bold=True, color=GREEN)

    title = lines[0].lstrip("# ").strip()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run(title)
    set_run_font(r, size=28, bold=True, color=DARK)

    rule = doc.add_paragraph()
    pPr = rule._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "24")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "06C755")
    pbdr.append(bottom)
    pPr.append(pbdr)

    for line in lines[1:]:
        clean = line.replace("**", "").strip()
        if not clean:
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(clean)
        if "期中專題報告" in clean:
            set_run_font(r, size=14, bold=True, color=GREEN)
        else:
            set_run_font(r, size=10.5, color=MUTED)
    doc.add_page_break()


def convert(md_path: Path, out_path: Path, cover=False):
    doc = Document()
    setup(doc)
    lines = md_path.read_text(encoding="utf-8").splitlines()

    if cover:
        split = lines.index("---")
        add_cover(doc, lines[:split])
        lines = lines[split + 1:]

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue

        if line.startswith("```"):
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            add_code(doc, code)
            i += 1
            continue

        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            heading(doc, line.lstrip("# ").strip(), level)
            i += 1
            continue

        if line.startswith("|") and i + 1 < len(lines) and set(lines[i + 1].replace("|", "").strip()) <= set("-: "):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                raw = lines[i].strip().strip("|")
                if not set(raw.replace("|", "").strip()) <= set("-: "):
                    rows.append([cell.strip() for cell in raw.split("|")])
                i += 1
            if rows:
                add_table(doc, rows)
            continue

        if line.startswith("!["):
            m = re.match(r"!\[[^\]]*\]\(([^)]+)\)", line)
            if m:
                add_image(doc, md_path.parent / m.group(1))
            i += 1
            continue

        if line.startswith(">"):
            quote = []
            while i < len(lines) and lines[i].startswith(">"):
                quote.append(lines[i])
                i += 1
            add_quote(doc, quote)
            continue

        if line.startswith("- [ ]"):
            p = body_para(doc, "☐ " + line[5:].strip(), style="List Bullet")
            i += 1
            continue

        if line.startswith("- "):
            body_para(doc, line[2:].strip(), style="List Bullet")
            i += 1
            continue

        if re.match(r"^\d+\.\s", line):
            body_para(doc, re.sub(r"^\d+\.\s*", "", line), style="List Number")
            i += 1
            continue

        if line.strip() == "---":
            rule = doc.add_paragraph()
            rule.paragraph_format.space_before = Pt(4)
            rule.paragraph_format.space_after = Pt(6)
            pPr = rule._p.get_or_add_pPr()
            pbdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "4")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "DBE7DE")
            pbdr.append(bottom)
            pPr.append(pbdr)
            i += 1
            continue

        para = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(
            r"^(#|\||- |\d+\.\s|>|```|!\[|---$)", lines[i].strip()
        ):
            para.append(lines[i].strip())
            i += 1
        body_para(doc, " ".join(para))

    doc.save(out_path)
    print(out_path)


if __name__ == "__main__":
    convert(DOCS / "專題報告.md", DOCS / "專題報告.docx", cover=True)
    convert(DOCS / "AI對話摘要.md", DOCS / "AI對話摘要.docx", cover=False)
